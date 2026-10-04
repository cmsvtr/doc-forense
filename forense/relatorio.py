"""Relatório Word de apoio ao analista (guia de leitura para a nota técnica, não é anexo)."""

import os
import re
from collections import defaultdict
from pathlib import Path
from typing import NamedTuple
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import __version__
from . import termos_cartel as T
from .caso import Caso, agora, ler_json
from .regex_br import extrair_datas
from .sei import localizador
from .texto import normalizar, trecho
from .triagem import ler_triagem

MAX_PRIORIDADES = 25
MAX_DATAS_CONTEXTO = 300

_FORTES = [re.compile(p) for termos in T.CATEGORIAS.values() for p, peso in termos if peso >= 2]


# ------------------------------------------------------------------ links para os originais

class Link(NamedTuple):
    texto: str
    alvo: str  # caminho relativo ao relatório, codificado como URL


def _alvo(caso: Caso, caminho_no_caso: str) -> str:
    """Caminho do original relativo à pasta relatorios/ («../originais/…»).

    Relativo, e não absoluto, para o link continuar funcionando se a pasta do caso inteira for
    movida ou copiada. O caminho real do arquivo é usado, inclusive dentro das subpastas de anexo.
    """
    rel = os.path.relpath(caso.raiz / caminho_no_caso, caso.relatorios).replace(os.sep, "/")
    return quote(rel, safe="/.-_()")


def _link(paragrafo, texto: str, alvo: str, negrito: bool = False, tamanho: float | None = None):
    """Acrescenta ao parágrafo um hiperlink do Word para o arquivo."""
    r_id = paragrafo.part.relate_to(alvo, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    h.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if negrito:
        rpr.append(OxmlElement("w:b"))
    cor = OxmlElement("w:color")
    cor.set(qn("w:val"), "0563C1")
    rpr.append(cor)
    if tamanho:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(tamanho * 2)))
        rpr.append(sz)
    sub = OxmlElement("w:u")
    sub.set(qn("w:val"), "single")
    rpr.append(sub)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    h.append(run)
    paragrafo._p.append(h)


# ------------------------------------------------------------------ helpers de formatação

def _sombrear(celula, cor_hex: str):
    tc_pr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    tc_pr.append(shd)


def _tabela(doc, cabecalho: list[str], linhas: list[list], larguras_cm: list[float] | None = None, fonte=8):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(cabecalho):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(fonte)
        _sombrear(c, "D9E2F3")
    for linha in linhas:
        cells = t.add_row().cells
        for i, v in enumerate(linha):
            cells[i].text = ""
            if isinstance(v, Link):
                _link(cells[i].paragraphs[0], v.texto, v.alvo, tamanho=fonte)
                continue
            r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(fonte)
    if larguras_cm:
        for row in t.rows:
            for i, w in enumerate(larguras_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def _aviso(doc, texto: str):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    _sombrear(c, "FFF2CC")
    c.text = ""
    r = c.paragraphs[0].add_run(texto)
    r.font.size = Pt(9)
    doc.add_paragraph()


def _paginas(lista) -> str:
    return ", ".join(map(str, lista))


# ------------------------------------------------------------------ seções

def _visao_geral(doc, docs, manifesto):
    total_pag = sum(len(d["paginas"]) for d in docs)
    ocr = sum(1 for d in docs for p in d["paginas"] if p["metodo"] == "ocr")
    baixa = sum(1 for d in docs for p in d["paginas"] if p.get("confianca_baixa"))
    erros = sum(1 for d in docs if d["status"] != "ok")
    datas_email = sorted(m["data_iso"][:10] for d in docs for m in d.get("emails", []) if m.get("data_iso"))
    doc.add_heading("1. Visão geral", level=1)
    itens = [
        f"Documentos distintos: {len(docs)} (PDF: {sum(1 for d in docs if d['arquivo']['tipo'] == 'pdf')}, "
        f"HTML: {sum(1 for d in docs if d['arquivo']['tipo'] == 'html')})",
        f"Páginas: {total_pag}, das quais {ocr} lidas por OCR ({baixa} com confiança baixa)",
        f"Documentos com erro ou extração parcial: {erros}",
        f"Mensagens de e-mail identificadas: {sum(len(d.get('emails', [])) for d in docs)}",
    ]
    if datas_email:
        itens.append(f"Período coberto pelos e-mails: {datas_email[0]} a {datas_email[-1]}")
    if manifesto and manifesto.get("arquivos_ignorados"):
        itens.append(f"Arquivos ignorados por formato não suportado: {len(manifesto['arquivos_ignorados'])} (ver seção 3)")
    for i in itens:
        doc.add_paragraph(i, style="List Bullet")


def _prioridades(doc, triagem, caso: Caso):
    doc.add_heading("2. Por onde começar: prioridades de leitura", level=1)
    doc.add_paragraph(
        "Primeiro os e-mails, os contratos e os documentos que citam 2 ou mais empresas ou 2 ou mais pessoas; "
        "depois os demais. Em cada grupo, a ordem é pela pontuação, que soma termos típicos de condutas "
        "colusivas (com peso) e sinais estruturais, como vários CNPJs ou domínios de e-mail no mesmo documento. "
        "Serve para decidir a ordem de leitura, não para concluir nada. Documento com pontuação baixa pode ser "
        "decisivo (por exemplo, linguagem cifrada).")
    relevantes = [r for r in triagem["documentos"] if r["pontuacao"] > 0 or r.get("prioritario")][:MAX_PRIORIDADES]
    if not relevantes:
        doc.add_paragraph("Nenhum documento pontuou na triagem.")
        return
    for r in relevantes:
        h = doc.add_heading(f"#{r['posicao']} — {r['arquivo']}  ({r['pontuacao']} pontos)", level=2)
        h.runs[0].font.size = Pt(11)
        alvo = _alvo(caso, r["caminho"])
        p = doc.add_paragraph()
        p.add_run("Citar como: ").bold = True
        p.add_run(r.get("localizador") or r["arquivo"])
        p.add_run("  ·  Arquivo: ").bold = True
        _link(p, r["caminho"], alvo)
        if r.get("motivos_prioridade"):
            p = doc.add_paragraph()
            p.add_run("Prioridade: ").bold = True
            p.add_run(", ".join(r["motivos_prioridade"]))
        if r["bonus"]:
            p = doc.add_paragraph()
            p.add_run("Sinais estruturais: ").bold = True
            p.add_run("; ".join(r["bonus"]))
        for cat, achados in r["categorias"].items():
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{cat}: ").bold = True
            p.add_run("; ".join(
                f"“{a['trechos'][0]['termo'] if a['trechos'] else a['padrao']}” ×{a['ocorrencias']} (p. {_paginas(a['paginas'])})"
                for a in achados))
        for t in _trechos_distintos(r, maximo=3):
            q = doc.add_paragraph(style="Quote")
            _link(q, f"p. {t['pagina']}", alvo, negrito=True)
            q.add_run(": " + t["trecho"])


def _trechos_distintos(r: dict, maximo: int) -> list[dict]:
    """Trechos de termos fortes, sem repetir a mesma região do texto."""
    escolhidos: list[dict] = []
    candidatos = [t for achados in r["categorias"].values() for a in achados if a["peso"] >= 2 for t in a["trechos"]]
    for t in candidatos:
        if any(t["pagina"] == e["pagina"] and abs(t.get("posicao", 0) - e.get("posicao", 0)) < 240 for e in escolhidos):
            continue
        escolhidos.append(t)
        if len(escolhidos) == maximo:
            break
    return escolhidos


def _qualidade(doc, docs, manifesto):
    doc.add_heading("3. Pontos de atenção na extração", level=1)
    doc.add_paragraph(
        "Antes de citar um trecho lido por OCR, confira-o na imagem original, sobretudo nas páginas abaixo. "
        "Erros de OCR em números (valores, datas, CNPJs) são comuns.")
    linhas = []
    for d in docs:
        for a in d["alertas"]:
            linhas.append([d["arquivo"]["nome"], a])
        for e in d["erros"]:
            if "Traceback" not in e:
                linhas.append([d["arquivo"]["nome"], f"Erro: {e}"])
    for d in docs:
        if not d.get("sei") and re.match(r"^\s*doc(umento)?\.?\s*n?[º°o.]?\s*\d", Path(d["arquivo"]["nome"]).stem, re.I):
            linhas.append([d["arquivo"]["nome"], "Sem número SEI: o arquivo está fora da pasta de anexo que dá o número "
                                                 f"({d['arquivo']['caminhos'][0]}). Copie a pasta dos autos com as subpastas."])
    duplicados = [d for d in docs if len(d["arquivo"]["caminhos"]) > 1]
    for d in duplicados:
        linhas.append([d["arquivo"]["nome"], "Arquivo idêntico (mesmo hash) em: " + "; ".join(d["arquivo"]["caminhos"])])
    for ign in (manifesto or {}).get("arquivos_ignorados", []):
        linhas.append([Path(ign).name, "Formato não suportado; não foi lido (verificar manualmente)."])
    if linhas:
        _tabela(doc, ["Documento", "Ponto de atenção"], linhas, [6, 11])
    else:
        doc.add_paragraph("Nenhum alerta.")


def _linha_do_tempo(doc, docs, caso: Caso):
    doc.add_heading("4. Linha do tempo preliminar (sem IA)", level=1)
    doc.add_paragraph(
        "Construída só com dados objetivos: cabeçalhos de e-mail e datas citadas no texto. "
        "A interpretação dos eventos fica para a etapa de análise.")

    doc.add_heading("4.1 Mensagens de e-mail", level=2)
    msgs = sorted(
        ([m.get("data_iso") or "", m.get("de", ""), m.get("para", ""), m.get("assunto", ""),
          Link(localizador(d), _alvo(caso, d["arquivo"]["caminhos"][0]))]
         for d in docs for m in d.get("emails", [])),
        key=lambda x: x[0] or "9999")
    if msgs:
        _tabela(doc, ["Data", "De", "Para", "Assunto", "Documento"],
                [[a.replace("T", " "), b, c, s, f] for a, b, c, s, f in msgs], [2.6, 3.6, 3.6, 4, 3.2])
    else:
        doc.add_paragraph("Nenhum cabeçalho de e-mail identificado.")

    doc.add_heading("4.2 Datas citadas perto de termos relevantes", level=2)
    doc.add_paragraph("Datas que aparecem a até ~300 caracteres de um termo de peso 2 ou mais da triagem.")
    eventos = []
    for d in docs:
        for pg in d["paginas"]:
            original = pg.get("texto") or ""
            norm = normalizar(original)
            for dt in extrair_datas(original):
                janela = norm[max(0, dt["inicio"] - 300): dt["fim"] + 300]
                if any(r.search(janela) for r in _FORTES):
                    eventos.append([dt["data"], Link(localizador(d, [pg["n"]]), _alvo(caso, d["arquivo"]["caminhos"][0])),
                                    trecho(original, dt["inicio"], dt["fim"], 160)])
    eventos.sort(key=lambda e: e[0])
    if eventos:
        if len(eventos) > MAX_DATAS_CONTEXTO:
            doc.add_paragraph(f"Exibindo {MAX_DATAS_CONTEXTO} de {len(eventos)}; use a busca do aplicativo para o restante.")
        _tabela(doc, ["Data", "Onde", "Contexto"], eventos[:MAX_DATAS_CONTEXTO], [2.2, 4.3, 10.5])
    else:
        doc.add_paragraph("Nenhuma data próxima de termos relevantes.")


def _entidades(doc, docs):
    doc.add_heading("5. Identificadores encontrados (CNPJ, CPF, e-mail, valores)", level=1)
    doc.add_paragraph(
        "Extraídos por padrão com validação de dígito verificador (CNPJ numérico e alfanumérico, CPF). "
        "Ponto de partida para o dramatis personae: confirme a razão social de cada CNPJ em fonte oficial.")
    por_tipo: dict[str, dict[str, list]] = defaultdict(lambda: defaultdict(list))
    for d in docs:
        for tipo, itens in (d.get("entidades") or {}).items():
            if tipo == "datas":
                continue
            for it in itens:
                por_tipo[tipo][it["valor"]].append(localizador(d, it["paginas"]))
    rotulos = {"cnpjs": "5.1 CNPJs", "cpfs": "5.2 CPFs", "emails": "5.3 E-mails", "valores": "5.4 Valores monetários"}
    for tipo, rotulo in rotulos.items():
        doc.add_heading(rotulo, level=2)
        itens = sorted(por_tipo.get(tipo, {}).items(), key=lambda kv: (-len(kv[1]), kv[0]))
        if not itens:
            doc.add_paragraph("Nenhum.")
            continue
        _tabela(doc, ["Valor", "Nº docs", "Onde aparece"],
                [[v, len(onde), "; ".join(onde[:8]) + (" …" if len(onde) > 8 else "")] for v, onde in itens[:200]],
                [4.5, 1.5, 11])


def _comunicacoes(doc, caso: Caso):
    from .comunicacoes import construir_teia

    t = construir_teia(caso)
    if not t["pares"]:
        return
    doc.add_heading("6. Comunicações entre pessoas (sem IA)", level=1)
    doc.add_paragraph(
        f"A partir de cabeçalhos de e-mail (De → Para e Cc) e de conversas exportadas: {t['mensagens']} mensagem(ns), "
        f"{len(t['pares'])} par(es) de pessoas. A pessoa é identificada pelo endereço de e-mail quando há; "
        "organização é o domínio do e-mail. Confira cada par nos documentos indicados.")
    if t["entre_organizacoes"]:
        doc.add_heading("6.1 Entre organizações diferentes", level=2)
        _tabela(doc, ["Organizações", "Mensagens", "Pares de pessoas"],
                [[" ↔ ".join(o["organizacoes"]), o["total"], o["pares"]] for o in t["entre_organizacoes"][:30]],
                [9, 3, 3])
    doc.add_heading("6.2 Pares que mais se comunicam", level=2)
    linhas = []
    for p in t["pares"][:40]:
        primeira = p["mensagens"][0]
        linhas.append([p["rotulo_a"] + (f" ({p['organizacao_a']})" if p["organizacao_a"] else ""),
                       p["rotulo_b"] + (f" ({p['organizacao_b']})" if p["organizacao_b"] else ""),
                       p["total"], f"{p['primeira'][:10]} a {p['ultima'][:10]}".strip(" a"),
                       Link("; ".join(p["documentos"][:3]) + (" …" if len(p["documentos"]) > 3 else ""),
                            _alvo(caso, primeira["caminho"]))])
    _tabela(doc, ["Pessoa A", "Pessoa B", "Mensagens", "Período", "Documentos"], linhas, [4, 4, 1.8, 3, 4.2])


def _validados(doc, caso: Caso):
    """Seções da etapa 2: só o que o analista validou."""
    from .analise_ia import achados_do_caso
    from .consolidacao import dramatis_personae, linha_do_tempo

    achados = achados_do_caso(caso)
    if not achados:
        return
    n_val = sum(1 for a in achados if a["revisao"]["status"] == "validado")
    n_pend = sum(1 for a in achados if a["revisao"]["status"] == "pendente")
    doc.add_heading("7. Dramatis personae (validado pelo analista)", level=1)
    doc.add_paragraph(
        f"Proposto pela IA local, conferido pela máquina (cada trecho existe no documento) e validado pelo analista. "
        f"{n_val} achado(s) validado(s); {n_pend} ainda pendente(s), fora deste relatório.")
    dp = dramatis_personae(achados)
    for rotulo, itens, campos in (("7.1 Pessoas", dp["pessoas"], ("cargos", "empresas")),
                                  ("7.2 Empresas", dp["empresas"], ("cnpjs",))):
        doc.add_heading(rotulo, level=2)
        if not itens:
            doc.add_paragraph("Nenhum item validado.")
            continue
        for i in itens:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(i["nome"]).bold = True
            extra = "; ".join(", ".join(i[c]) for c in campos if i[c])
            if extra:
                p.add_run(f" — {extra}")
            for f in i["fontes"][:3]:
                q = doc.add_paragraph(style="Quote")
                _link(q, f["localizador"], _alvo(caso, f["caminho"]), negrito=True)
                q.add_run(": «" + " ".join(f["trecho"].split()) + "»")

    doc.add_heading("8. Linha do tempo (validada pelo analista)", level=1)
    lt = linha_do_tempo(achados)
    if not lt:
        doc.add_paragraph("Nenhum evento validado.")
    else:
        doc.add_paragraph("A coluna «Resumo» é da IA e não é citação; o trecho é o texto do documento.")
        _tabela(doc, ["Data", "Categoria", "Resumo (IA)", "Trecho do documento", "Onde"],
                [[e["data"] or "sem data", e["categoria"], e["descricao_ia"], "«" + " ".join(e["trecho"].split()) + "»",
                  Link(e["localizador"], _alvo(caso, e["caminho"]))] for e in lt],
                [2, 3, 3.5, 6, 2.5])


def _anexo_custodia(doc, docs, caso: Caso, manifesto):
    doc.add_page_break()
    doc.add_heading("Anexo A — Inventário e cadeia de custódia", level=1)
    hash_manifesto = None
    arq = caso.raiz / "manifesto.json.sha256"
    if arq.exists():
        hash_manifesto = arq.read_text(encoding="utf-8").split()[0]
    doc.add_paragraph(
        f"Manifesto: {caso.manifesto.name}" + (f" — SHA-256 {hash_manifesto}" if hash_manifesto else "")
        + (f" — gerado em {manifesto['gerado_em']}" if manifesto else ""))
    linhas = []
    for i, d in enumerate(docs, 1):
        ocr = sum(1 for p in d["paginas"] if p["metodo"] == "ocr")
        sei = (d.get("sei") or {}).get("numero") or "—"
        linhas.append([i, Link(d["arquivo"]["caminhos"][0], _alvo(caso, d["arquivo"]["caminhos"][0])), sei, d["arquivo"]["tipo"].upper(), len(d["paginas"]),
                       ocr, d["status"], d["arquivo"]["sha256"]])
    _tabela(doc, ["#", "Arquivo", "SEI", "Tipo", "Pág.", "OCR", "Status", "SHA-256"], linhas,
            [0.7, 4.6, 1.4, 1, 0.9, 0.9, 1.2, 6.3], fonte=7)


def _anexo_metodo(doc, manifesto, triagem):
    doc.add_heading("Anexo B — Método e ferramentas", level=1)
    if manifesto:
        f = manifesto.get("ferramentas", {})
        p = manifesto.get("parametros", {})
        for item in [
            f"Aplicativo doc-forense {manifesto['aplicativo']['versao']} (extrator {manifesto['aplicativo']['extrator']})",
            f"Sistema: {manifesto['sistema']['so']}, Python {manifesto['sistema']['python']}",
            f"OCR: {f.get('tesseract') or 'indisponível'}; idioma {f.get('idioma', p.get('idioma'))}; "
            f"modelo {Path(f.get('traineddata', '')).name or '-'} (SHA-256 {f.get('traineddata_sha256', '-')})",
            f"Renderização de PDF a {p.get('dpi')} dpi; página tratada como digital se tiver ao menos "
            f"{p.get('limiar_texto_digital')} letras/dígitos; alerta de OCR abaixo de {p.get('confianca_baixa')}% de confiança",
            "Bibliotecas: " + ", ".join(f"{k} {v}" for k, v in f.get("bibliotecas", {}).items()),
        ]:
            doc.add_paragraph(item, style="List Bullet")
    if triagem:
        doc.add_paragraph(
            f"Triagem gerada em {triagem['gerado_em']} com a lista de termos de hash {triagem['termos_hash']} "
            "(arquivo forense/termos_cartel.py). Pontos = soma, por termo, de min(ocorrências, "
            f"{T.MAX_OCORRENCIAS_POR_TERMO}) × peso, mais bônus estruturais.", style="List Bullet")
    doc.add_paragraph(
        "O texto extraído de cada documento está em extraido/<id>.json, página a página, com o método usado "
        "(texto digital ou OCR) e a confiança do OCR. Os originais não são alterados; a integridade pode ser "
        "verificada a qualquer momento pelo aplicativo (recalcula os hashes e compara com o manifesto).",
        style="List Bullet")


# ------------------------------------------------------------------ principal

def gerar_relatorio(caso: Caso) -> Path:
    docs = caso.documentos()
    if not docs:
        raise RuntimeError("Nenhum documento processado neste caso.")
    triagem = ler_triagem(caso) or {"documentos": [], "gerado_em": "-", "termos_hash": "-"}
    manifesto = ler_json(caso.manifesto) if caso.manifesto.exists() else None

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(2))
    sec.top_margin = sec.bottom_margin = Cm(2)

    titulo = doc.add_heading("Relatório de apoio ao analista", level=0)
    titulo.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    doc.add_paragraph(f"Caso: {caso.nome}  ·  Gerado em {agora()}  ·  doc-forense {__version__}")
    _aviso(doc,
           "DOCUMENTO DE TRABALHO INTERNO. Organiza o material para orientar a leitura e a redação da nota técnica. "
           "Não é prova nem conclusão. Todo trecho deve ser conferido no documento original (arquivo e página indicados), "
           "especialmente os lidos por OCR. A triagem é heurística e pode deixar passar documentos relevantes.")

    doc.add_paragraph(
        "Os nomes de arquivo e as páginas em azul são links para o original (Ctrl+clique). O Word pode pedir "
        "confirmação antes de abrir. O link abre o arquivo, não a página: vá à página indicada. Os links "
        "funcionam enquanto este relatório estiver na pasta relatorios do caso.")

    _visao_geral(doc, docs, manifesto)
    _prioridades(doc, triagem, caso)
    _qualidade(doc, docs, manifesto)
    _linha_do_tempo(doc, docs, caso)
    _entidades(doc, docs)
    _comunicacoes(doc, caso)
    _validados(doc, caso)
    _anexo_custodia(doc, docs, caso, manifesto)
    _anexo_metodo(doc, manifesto, triagem)

    caso.relatorios.mkdir(parents=True, exist_ok=True)
    destino = caso.relatorios / f"relatorio_apoio_{agora()[:19].replace(':', '').replace('-', '')}.docx"
    doc.save(destino)
    caso.registrar("relatorio_gerado", arquivo=caso.relativo(destino))
    return destino
