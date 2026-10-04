"""Fichas individuais, geradas sob demanda para as pessoas escolhidas no dramatis personae.

Junta o que foi validado sobre a pessoa (achados da IA, com as correções do analista) com o que a
teia de comunicações sabe dela (endereços, organização, contrapartes, mensagens). Todo item guarda
o documento e a página de origem, para o botão de abrir, e o alvo, para o botão de corrigir.
"""

from pathlib import Path

from .citacao import aparece
from .consolidacao import chave_nome, dramatis_personae
from .nomes import chave


def item_do_achado(a: dict, descricao: str) -> dict:
    """Item exibível (ficha, linha do tempo) com o que os botões de abrir e de corrigir precisam."""
    return {"data": a["dados"].get("data", ""), "descricao": descricao, "localizador": a["localizador"],
            "caminho": a["caminho"], "pagina": a["pagina"], "trecho": a["trecho_fonte"], "origem": "achado",
            "tipo": a["tipo"], "alvo": {"tipo": "achado", "id": a["id"], "documento_id": a["documento_id"],
                                        "pagina": a["pagina"], "localizador": a["localizador"], "trecho": a["trecho_fonte"]},
            "dados": a["dados"], "correcoes": a.get("correcoes", {}), "status": a["revisao"]["status"],
            "contexto": a.get("ia") or {}}


def item_da_mensagem(m: dict, teia: dict) -> dict:
    de, para = teia["pessoas"][m["de"]]["rotulo"], teia["pessoas"][m["para"]]["rotulo"]
    return {"data": m["data"], "descricao": f"{m['meio']}: {de} → {para}" + (f" · «{m['assunto']}»" if m["assunto"] else ""),
            "localizador": m["localizador"], "caminho": m["caminho"], "pagina": m["pagina"], "trecho": "",
            "origem": "mensagem", "tipo": "mensagem",
            "alvo": {"tipo": "mensagem", "id": m["id"], "documento_id": m["documento_id"],
                     "pagina": m["pagina"], "localizador": m["localizador"], "trecho": ""},
            # de/para completos (todos os destinatários da mensagem, não só o deste par): a correção os substitui
            "dados": {"de": m.get("de_texto", de), "para": m.get("para_texto", para), "data": m["data"],
                      "assunto": m["assunto"]},
            "correcoes": m.get("correcoes", {}), "status": "máquina",
            "contexto": {"origem": "cabeçalho ou conversa, lido pela máquina"}}


def montar(pessoa: dict, achados: list[dict], teia: dict, incluir_pendentes: bool = False) -> dict:
    """pessoa: um item de dramatis_personae()["pessoas"]."""
    chaves = {chave_nome(g) for g in pessoa["grafias"]} | {chave_nome(pessoa["nome"])}
    aceitos = {"validado", "pendente"} if incluir_pendentes else {"validado"}
    validos = [a for a in achados if a["revisao"]["status"] in aceitos]

    # quem é a pessoa na teia: o nome casa com algum nome visto em cabeçalhos ou conversas
    ids_teia = {pid for pid, p in teia["pessoas"].items() if any(chave(n) in chaves for n in p["nomes"])}
    enderecos = sorted({teia["pessoas"][i]["endereco"] for i in ids_teia if teia["pessoas"][i]["endereco"]})
    organizacoes = sorted({teia["pessoas"][i]["organizacao"] for i in ids_teia if teia["pessoas"][i]["organizacao"]})

    contrapartes, linha = [], []
    for par in teia["pares"]:
        if par["a"] in ids_teia or par["b"] in ids_teia:
            outro = par["b"] if par["a"] in ids_teia else par["a"]
            contrapartes.append({"pessoa": teia["pessoas"][outro]["rotulo"], "organizacao": teia["pessoas"][outro]["organizacao"],
                                 "mensagens": par["total"], "primeira": par["primeira"], "ultima": par["ultima"],
                                 "entre_organizacoes": par["entre_organizacoes"]})
            linha += [item_da_mensagem(m, teia) for m in par["mensagens"]]
    contrapartes.sort(key=lambda c: -c["mensagens"])

    citacoes = []
    for a in validos:
        if a["tipo"] == "eventos":
            participa = any(chave_nome(p) in chaves for p in a["dados"].get("participantes", []))
            if participa or any(aparece(g, a["trecho_fonte"]) for g in pessoa["grafias"]):
                d = a["dados"]
                linha.append(item_do_achado(a, f"{d.get('categoria', '')}: {d.get('descricao_ia', '')}"))
        elif a["tipo"] == "pessoas" and chave_nome(a["dados"]["nome"]) in chaves:
            extra = ", ".join(x for x in (a["dados"].get("cargo"), a["dados"].get("empresa")) if x)
            citacoes.append(item_do_achado(a, a["dados"]["nome"] + (f" — {extra}" if extra else "")))

    # mensagens repetidas entre pares (a mesma mensagem com vários destinatários) aparecem uma vez por par
    linha.sort(key=lambda i: (i["data"] == "", i["data"] or "", i["localizador"]))
    return {"nome": pessoa["nome"], "grafias": pessoa["grafias"], "cargos": pessoa.get("cargos", []),
            "empresas": pessoa.get("empresas", []), "enderecos": enderecos, "organizacoes": organizacoes,
            "contrapartes": contrapartes, "linha_do_tempo": linha, "citacoes": citacoes,
            "documentos": sorted({i["localizador"].rsplit(", p.", 1)[0] for i in linha + citacoes})}


def fichas(caso, nomes: list[str], incluir_pendentes: bool = False) -> list[dict]:
    from .analise_ia import achados_do_caso
    from .comunicacoes import construir_teia

    achados = achados_do_caso(caso)
    teia = construir_teia(caso)
    dp = {p["nome"]: p for p in dramatis_personae(achados, incluir_pendentes=incluir_pendentes)["pessoas"]}
    return [montar(dp[n], achados, teia, incluir_pendentes) for n in nomes if n in dp]


def gerar_word(caso, ficha: dict) -> Path:
    """Word da ficha, sob demanda, com os links para os originais."""
    import re

    from docx import Document
    from docx.shared import Pt

    from .caso import agora
    from .relatorio import Link, _alvo, _aviso, _link, _tabela

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(f"Ficha individual — {ficha['nome']}", level=0)
    doc.add_paragraph(f"Caso: {caso.nome}  ·  Gerada em {agora()}")
    _aviso(doc, "DOCUMENTO DE TRABALHO INTERNO. Reúne o que o analista validou e o que a máquina extraiu de cabeçalhos "
                "de e-mail e conversas. Itens marcados com ✎ foram corrigidos pelo analista. Confira cada item na fonte.")
    doc.add_heading("Identificação", level=1)
    for rotulo, valor in (("Grafias", ficha["grafias"]), ("Cargos", ficha["cargos"]), ("Empresas", ficha["empresas"]),
                          ("E-mails", ficha["enderecos"]), ("Organizações (domínio)", ficha["organizacoes"])):
        if valor:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{rotulo}: ").bold = True
            p.add_run("; ".join(valor))
    if ficha["contrapartes"]:
        doc.add_heading("Com quem se comunicou", level=1)
        _tabela(doc, ["Pessoa", "Organização", "Mensagens", "Período"],
                [[c["pessoa"], c["organizacao"] or "—", c["mensagens"], f"{c['primeira'][:10]} a {c['ultima'][:10]}".strip(" a")]
                 for c in ficha["contrapartes"]], [5, 4, 2, 4])
    doc.add_heading("Linha do tempo", level=1)
    if ficha["linha_do_tempo"]:
        _tabela(doc, ["Data", "O quê", "Trecho", "Onde"],
                [[(i["data"] or "sem data").replace("T", " ")[:16], i["descricao"] + (" ✎" if i["correcoes"] else ""),
                  ("«" + " ".join(i["trecho"].split()) + "»") if i["trecho"] else "",
                  Link(i["localizador"], _alvo(caso, i["caminho"]))] for i in ficha["linha_do_tempo"]],
                [2.4, 5, 6, 3.6])
    else:
        doc.add_paragraph("Nenhum evento ou mensagem datada.")
    if ficha["citacoes"]:
        doc.add_heading("Onde é citada", level=1)
        for i in ficha["citacoes"]:
            q = doc.add_paragraph(style="Quote")
            _link(q, i["localizador"], _alvo(caso, i["caminho"]), negrito=True)
            q.add_run(f": {i['descricao']}{' ✎' if i['correcoes'] else ''} — «{' '.join(i['trecho'].split())}»")
    caso.relatorios.mkdir(parents=True, exist_ok=True)
    nome_arq = re.sub(r"[^\w\-]+", "_", ficha["nome"]).strip("_")[:60] or "pessoa"
    destino = caso.relatorios / f"ficha_{nome_arq}_{agora()[:19].replace(':', '').replace('-', '')}.docx"
    doc.save(destino)
    caso.registrar("ficha_gerada", pessoa=ficha["nome"], arquivo=caso.relativo(destino))
    return destino
